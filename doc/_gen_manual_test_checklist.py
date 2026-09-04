# -*- coding: utf-8 -*-
"""Generate manual functional test checklist Excel for PayloadGroundTest."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "平台人工功能测试清单.xlsx"

# (模块, 页面标题, 测试内容)
ROWS: list[tuple[str, str, str]] = []


def add(module: str, page: str, *items: str) -> None:
    for it in items:
        ROWS.append((module, page, it))


# ─── 登录与壳层 ─────────────────────────────────────────────
add(
    "登录与壳层",
    "登录页 /login",
    "打开登录页，页面标题、表单、验证码（若开启）显示正常",
    "错误账号/密码登录，提示明确且不进入系统",
    "正确账号密码（如 admin/admin123）登录成功，进入首页",
    "验证码错误时拒绝登录（验证码开启时）",
    "勾选「记住密码」后重新打开仍保留账号（若功能开启）",
    "未登录直接访问业务路由，跳转登录页",
)
add(
    "登录与壳层",
    "注册页 /register",
    "注册入口可见性符合系统开关配置（关闭时不可注册）",
    "注册开关开启时：合法信息可注册；非法/重复用户名有提示",
)
add(
    "登录与壳层",
    "顶栏 / 布局",
    "顶栏可切换主题模式、布局大小",
    "打开「布局设置」抽屉，修改菜单/主题等后界面即时生效",
    "「清除缓存」类设置项可恢复默认布局偏好",
    "Tags-View 多页签可打开、切换、关闭",
    "退出登录后会话失效，再访问需重新登录",
)
add(
    "登录与壳层",
    "个人中心 /user/profile",
    "查看并修改基本资料后保存成功，刷新仍生效",
    "修改密码：原密码错误拒绝；正确原密码可修改，新密码可登录",
    "上传/更换头像成功并显示",
)

# ─── 首页设备服务 ───────────────────────────────────────────
add(
    "首页",
    "设备服务 /index",
    "进入首页，设备服务列表加载；列含类型/设备ID/连接信息/来源/组装器/解释器/打开时间/操作",
    "「刷新」更新列表；开启「自动刷新」后列表周期性更新，关闭后停止",
    "「新建 CAN」：填写厂商/索引/通道/波特率等及组装器/解释器，打开成功并出现在列表",
    "「新建 UDP」：填写本机地址/端口及组装器/解释器，打开成功",
    "「新建串口」：选择端口与物理参数及组装器/解释器，打开成功；已占用串口不可重复打开",
    "行内「修改」：可改绑组装器（空=透传）、解释器（空=解绑），保存后列表更新",
    "行内「关闭连接」二次确认后关闭，列表中该连接消失",
    "「关闭所有连接」二次确认后清空全部连接",
    "无连接时列表空状态正常；非法参数打开连接有明确错误提示",
)

# ─── 遥控 BIU / XL（共用页，分族测）─────────────────────────
for family, label in (("BIU", "BIU"), ("XL", "XL")):
    path_ctrl = f"/telecontrol/{family.lower()}/control"
    path_cmd = f"/telecontrol/{family.lower()}/command"
    path_seq = f"/telecontrol/{family.lower()}/sequence"
    add(
        f"遥控-{label}",
        f"{label}·控制 {path_ctrl}",
        "页面加载，CAN 工具栏可见；可新建/关闭 CAN-A、CAN-B，并可选择「当前发送」通道",
        "选择遥测类型后「发送遥测请求」成功（依赖已打开且绑定正确的 CAN）",
        "「定时遥测」打开后持续请求，关闭后停止",
        "设置载荷时间（含使用系统当前时间）；起始时间设置/重置；时间偏差(ms)设置成功",
        "「定时同步广播」打开/关闭状态切换正确"
        + ("；XL 可勾选 GNSS 有效" if family == "XL" else ""),
        "HEX 文本框输入合法数据后："
        + (
            "「发送遥控指令」「发送广播」均可发出"
            if family == "BIU"
            else "「发送遥控指令」「姿控广播」「GNSS 广播」均可发出"
        ),
        ("「CAN重置」可发送" if family == "BIU" else "选择系统指令后「发送」成功（含广播勾选场景）"),
        "未选发送通道或通道已关闭时，发送类操作有明确提示",
    )
    add(
        f"遥控-{label}",
        f"{label}·遥控 {path_cmd}",
        "加载遥控配置树（目录→指令），树与族（BIU/XL）配置一致",
        "搜索框：按代号/名称过滤；空格 AND、| OR、! NOT、通配符等精简语法可用",
        "点击目录：中间列出该目录全部指令卡片；点击单条指令：仅显示该条",
        "改参数（数值/下拉/科学计数等）时参数 HEX 实时更新",
        "「预览组帧」生成完整帧 HEX；「发送指令」成功并出现在右侧发送历史",
        "右侧历史显示时间/指令名/HEX/状态；「清空」清空历史",
        "无权限或未开 CAN 时发送失败提示合理",
    )
    add(
        f"遥控-{label}",
        f"{label}·指令序列 {path_seq}",
        "列表展示序列名称、指令条数、备注、创建时间；搜索/重置可用",
        "「新增」可保存序列（名称 + 指令行 HEX/间隔；可从遥控指令树挑选）",
        "「修改」编辑已有序列并保存；校验非法 HEX/空名称",
        "「复制」生成带「-副本」草稿并可保存为新序列",
        "「导出」导出序列指令数据",
        "「执行」：选择目标设备→开始执行；进度可见；可「后台继续」",
        "「日志」查看执行历史；「详情」查看单次发送明细",
        "单条/批量「删除」需确认，删除后列表更新",
        "进入编辑页（若跳转）：左侧指令树、中间预览组帧/设置指令、右侧排序增删移；未保存离开有提示；「保存」成功",
    )

# ─── 遥测 ───────────────────────────────────────────────────
add(
    "遥测",
    "实时数据 /telemetry/live",
    "表类型下拉含 BIU/XL/单板/相机等分组；切换表后表格字段与配置一致",
    "有实时数据时表格轮询刷新「当前值」；无数据时空状态正常",
    "刷新页面后仍保持上次选择的表类型（本地偏好）",
    "双击某「当前值」跳转实时曲线，并带上对应表/量",
)
add(
    "遥测",
    "实时曲线 /telemetry/curve",
    "可添加多条曲线（上限约 10 条），超出有提示",
    "曲线实时刷新；支持缩放、裁剪、重置时间窗、Y 轴自适应",
    "可移除单条曲线；「导出 CSV」文件内容与曲线数据对应",
)
add(
    "遥测",
    "历史CAN数据 /telemetry/canHistory",
    "选择遥测表 + 时间范围后「解析/打开」成功建立回放会话",
    "回放条：上一帧/下一帧、滑块跳转、自动播放与间隔设置工作正常",
    "帧表格展示归档数据（非 Redis 实时）；切帧后表格内容变化",
    "无数据/非法时间范围有明确提示",
)
add(
    "遥测",
    "历史CAN曲线 /telemetry/archive",
    "选择表/量与时间范围查询后绘制历史曲线",
    "支持多曲线、裁剪、重置/自适应、导出 CSV",
)
add(
    "遥测",
    "历史文件数据 /telemetry/fileHistory",
    "可浏览/选择服务端日志或上传本地日志文件",
    "「解析」后状态变为 ready；回放条与帧表可浏览各帧",
    "解析失败（坏文件）有错误提示",
)
add(
    "遥测",
    "历史文件曲线 /telemetry/fileCurve",
    "对已解析文件选量绘图；查询/重置/自适应/裁剪/CSV 导出可用",
)

# ─── 单板 ───────────────────────────────────────────────────
for ver, path, res_note in (
    ("v1.6", "/board/camera", "分辨率下拉 400/256/128/64"),
    ("v1.7", "/board/camera_v17", "边长输入 8–400 步进 8"),
):
    add(
        "单板",
        f"相机{ver} {path}",
        "「新建/关闭控制串口」「新建/关闭图像串口」成功；两路可独立开关",
        "左侧遥控树：预览组帧 / 发送指令 / 导出预览 HEX 可用",
        f"嵌入遥测表刷新正常；分辨率同步逻辑合理（{res_note}；手选后不被遥测覆盖）",
        "图像：单次刷新、连续刷新、停止刷新；勾选「首帧后自动拍照」「显示质心」生效",
        "「图片保存」「图片上传」可用；传输信息面板可切换多源 IO",
        "控制口/图像口未连接时相关操作有提示",
    )

add(
    "单板",
    "热控电机 /board/rkdj",
    "新建/关闭串口连接成功；传输信息显示收发",
    "指令区：预览组帧、发送、导出可用",
    "嵌入 RKDJ 遥测表轮询显示正常",
)
add(
    "单板",
    "CPA-ZK /board/zk",
    "新建/关闭串口连接成功；传输信息显示收发",
    "指令区：预览组帧、发送、导出可用",
    "嵌入 ZK 遥测表轮询显示正常",
)
add(
    "单板",
    "地检板 /board/dj",
    "新建/关闭 UDP 连接成功（地检工程遥测）",
    "指令区：预览组帧、发送、导出可用",
    "嵌入地检工程遥测表轮询显示正常",
)

# ─── LVDS / 调试 / 重构 ─────────────────────────────────────
add(
    "LVDS",
    "工程遥测 /lvds/engineering",
    "顶部状态：采样率/FPS/点数显示；「暂停/继续」切换数据采集或刷新",
    "左侧信号多选（最多 8 路），超出提示；「清空」取消全部选择",
    "中间多通道波形随选中信号更新",
    "右侧游标：C1/C2 开关、垂直/水平模式；ΔT/ΔY/Freq 计算合理",
)
add(
    "调试",
    "数据模拟 /debug/simulate",
    "选择组装器+解释器，填 HEX 后发送，写 Redis 来源为开发注入；实时遥测可见结果",
    "「示例数据」拉取黄金样本并填入 HEX（含相机 v16/v17 相关样例）",
    "清空输入区工作正常",
    "CAN 复合帧注入：填入样例、开始/停止模拟流程可用",
)
add(
    "调试",
    "数据收发 /debug/xfer",
    "设备列表可刷新；可选在线设备发原始 HEX（CAN 帧 ID+≤8 字节 / UDP 远端）",
    "IO 日志实时/刷新显示收发内容",
    "离线/历史设备仅可查看、不可发送（或有明确限制提示）",
)
add(
    "调试",
    "配置文件 /debug/configFile",
    "列出 TeleControl / TeleMetry 等配置 JSON",
    "单文件：预览、下载、编辑保存、单文件重载成功",
    "「热重载全部」后遥控/遥测页配置生效",
    "「导出指令」生成可用导出物",
    "无编辑权限时仅可查看（若权限分离）",
)
add(
    "调试",
    "遥测计算 /debug/tmCalc",
    "选择遥测表+字段，输入字段 HEX 后「计算」，结果含名称/当前值/单位/HEX",
    "计算结果插入历史表首行；「清空历史」清空",
)
add(
    "重构",
    "重构 /refactor",
    "页面可打开，显示建设中占位（确认无报错白屏）",
)

# ─── 系统管理 ───────────────────────────────────────────────
add(
    "系统管理",
    "用户管理 /system/user",
    "列表查询/重置；新增用户成功；修改用户信息成功",
    "删除/批量删除需确认；重置密码可用",
    "导入用户、导出用户文件可用",
    "「分配角色」勾选角色保存后，该用户权限符合角色",
    "停用用户后无法登录；启用后可登录",
)
add(
    "系统管理",
    "角色管理 /system/role",
    "角色增删改查、导出可用",
    "菜单权限勾选保存后，对应用户菜单/按钮可见性正确",
    "「分配用户」可绑定/取消用户",
)
add(
    "系统管理",
    "菜单管理 /system/menu",
    "菜单树展示正确；新增目录/菜单/按钮；修改、删除（有子节点时限制合理）",
)
add(
    "系统管理",
    "部门管理 /system/dept",
    "部门树增删改；用户归属部门可选到新建部门",
)
add(
    "系统管理",
    "岗位管理 /system/post",
    "岗位增删改查、导出；用户可关联岗位",
)
add(
    "系统管理",
    "字典管理 /system/dict",
    "字典类型增删改查、导出",
    "进入字典数据页：数据项增删改查，前端下拉引用生效",
)
add(
    "系统管理",
    "参数设置 /system/config",
    "参数增删改查、导出；修改后业务读取到新值（按缓存刷新规则）",
)
add(
    "系统管理",
    "通知公告 /system/notice",
    "公告增删改查；发布后相关展示位可见（若前端有展示）",
)
add(
    "系统管理",
    "操作日志 /system/log/operlog",
    "按条件查询操作日志；删除/清空；导出",
)
add(
    "系统管理",
    "登录日志 /system/log/logininfor",
    "查询登录成功/失败记录；删除/导出",
    "账户解锁（若有锁定机制）后可重新登录",
)

# ─── 系统监控 / 工具 ───────────────────────────────────────
add(
    "系统监控",
    "在线用户 /monitor/online",
    "列表显示当前在线会话；强制下线单用户/批量后对方需重新登录",
)
add(
    "系统监控",
    "定时任务 /monitor/job",
    "任务列表含「遥测归档月分区维护」等；新增/修改/删除/导出",
    "启用/暂停任务；「执行一次」立即跑通",
    "调度日志页可按任务查看执行历史",
)
add(
    "系统监控",
    "服务监控 /monitor/server",
    "展示 CPU/内存/JVM(或进程)/磁盘等信息，刷新正常",
)
add(
    "系统监控",
    "缓存监控 /monitor/cache",
    "展示 Redis 基本信息与统计",
)
add(
    "系统监控",
    "缓存列表 /monitor/cacheList",
    "可浏览缓存键；清理指定键或模式后数据消失",
)
add(
    "系统监控",
    "传输加密 /monitor/transportCrypto",
    "页面可打开并展示传输加密相关监控/状态（无白屏报错）",
)
add(
    "系统工具",
    "表单构建 /tool/build",
    "拖拽组件生成表单预览正常（工具页冒烟）",
)
add(
    "系统工具",
    "代码生成 /tool/gen",
    "导入表、预览代码、生成/下载；编辑生成配置页可保存",
)
add(
    "系统工具",
    "系统接口 /tool/swagger",
    "Swagger/OpenAPI 文档页可打开，抽样接口可试调（冒烟）",
)

# ─── 跨页 / 权限 / 稳定性 ───────────────────────────────────
add(
    "跨页与权限",
    "权限与会话",
    "无业务权限角色：对应菜单隐藏，直接 URL 访问被拦截或提示无权限",
    "有菜单无按钮权限：页面可进但增删改按钮不可见/不可用",
    "Token 过期或被强制下线后，操作提示重新登录",
)
add(
    "跨页与权限",
    "端到端冒烟（建议联调环境）",
    "首页开 CAN → BIU/XL 控制发遥测请求 → 实时数据/曲线可见对应字段变化",
    "遥控页发指令 → 发送历史有记录；指令序列执行整链成功",
    "调试·数据模拟注入 → 实时遥测表出现解析结果",
    "相机控制串口发指令 + 图像串口收图 → 图像区出图并可保存",
    "历史 CAN / 历史文件回放与曲线在有样例数据时可走通",
)
add(
    "跨页与权限",
    "异常与体验",
    "后端停止时前端请求失败提示友好，不整页崩溃",
    "浏览器缩小/常见分辨率下主要业务页布局可用（无严重遮挡）",
    "连续开关连接、反复进出会话页无内存泄漏级卡死（主观观察）",
)


def build() -> None:
    wb = Workbook()

    # —— 说明页 ——
    ws0 = wb.active
    ws0.title = "使用说明"
    ws0["A1"] = "载荷地检平台 — 人工功能测试清单"
    ws0["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="1F4E79")
    notes = [
        "",
        "用途：测试人员按「模块 → 页面」逐项执行，在「测试结果」列填写结论。",
        "结果取值：通过 / 不通过 / 阻塞 / 未测 / 不适用。",
        "粒度：覆盖各页面主要功能点；同一操作的细碎点击不拆分；BIU/XL 同源页面分族各测一遍。",
        "建议环境：前端 http://localhost:80（或实际部署地址）；后端与 Redis/DB 就绪；硬件或模拟器按需接入。",
        "默认账号：admin / admin123（以实际环境为准）。",
        "测试顺序建议：登录壳层 → 首页设备连接 → 遥控 BIU/XL → 实时遥测 → 指令序列 → 单板/相机 → 历史回放 → 调试 → LVDS → 系统管理/监控回归。",
        "「重构」页当前为占位，确认可打开即可。",
        "依据：菜单与页面实现（ruoyi-fastapi-frontend）、doc/05-前端页面设计.md。",
        "",
        "工作表说明：",
        "  · 功能测试清单 — 主表，逐项勾选",
        "  · 按模块汇总 — 各模块用例数，便于分工",
    ]
    for i, line in enumerate(notes, start=2):
        ws0[f"A{i}"] = line
        ws0[f"A{i}"].font = Font(name="微软雅黑", size=11)
    ws0.column_dimensions["A"].width = 110

    # —— 主清单 ——
    ws = wb.create_sheet("功能测试清单", 1)
    headers = [
        "序号",
        "模块",
        "页面标题",
        "测试内容",
        "测试结果",
        "缺陷编号/备注",
        "测试人",
        "测试日期",
    ]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    alt_fill = PatternFill("solid", fgColor="F2F7FB")
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    module_fills = {
        "登录与壳层": "E8F5E9",
        "首页": "FFF3E0",
        "遥控-BIU": "E3F2FD",
        "遥控-XL": "E8EAF6",
        "遥测": "F3E5F5",
        "单板": "E0F7FA",
        "LVDS": "FCE4EC",
        "调试": "FFF8E1",
        "重构": "ECEFF1",
        "系统管理": "EFEBE9",
        "系统监控": "E8EAF6",
        "系统工具": "E0F2F1",
        "跨页与权限": "FFEBEE",
    }

    for i, (module, page, content) in enumerate(ROWS, start=1):
        r = i + 1
        values = [i, module, page, content, "", "", "", ""]
        fill = PatternFill("solid", fgColor=module_fills.get(module, "FFFFFF"))
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = thin
            cell.alignment = center if c in (1, 5, 7, 8) else wrap
            if c == 2:
                cell.fill = fill
            elif c == 5:
                cell.fill = PatternFill("solid", fgColor="FFFDE7")
            elif i % 2 == 0 and c not in (2, 5):
                cell.fill = alt_fill

    widths = [8, 14, 36, 72, 12, 22, 10, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(ROWS) + 1}"
    ws.row_dimensions[1].height = 22
    for r in range(2, len(ROWS) + 2):
        ws.row_dimensions[r].height = 32

    dv = DataValidation(
        type="list",
        formula1='"通过,不通过,阻塞,未测,不适用"',
        allow_blank=True,
    )
    dv.error = "请选择列表中的结果"
    dv.errorTitle = "测试结果"
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(ROWS) + 1}")

    # —— 汇总 ——
    ws2 = wb.create_sheet("按模块汇总", 2)
    for col, h in enumerate(["模块", "用例数", "建议测试人", "完成数", "通过数", "备注"], 1):
        cell = ws2.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    from collections import Counter

    counts = Counter(m for m, _, _ in ROWS)
    for i, (mod, cnt) in enumerate(counts.items(), start=2):
        ws2.cell(i, 1, mod).border = thin
        ws2.cell(i, 2, cnt).border = thin
        for c in range(3, 7):
            ws2.cell(i, c, "").border = thin
        ws2.cell(i, 1).fill = PatternFill("solid", fgColor=module_fills.get(mod, "FFFFFF"))
    total_row = len(counts) + 2
    ws2.cell(total_row, 1, "合计").font = Font(name="微软雅黑", bold=True)
    ws2.cell(total_row, 2, len(ROWS)).font = Font(name="微软雅黑", bold=True)
    for c in range(1, 7):
        ws2.cell(total_row, c).border = thin
    for i, w in enumerate([16, 10, 14, 10, 10, 30], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT)
    print(f"Wrote {OUT}  rows={len(ROWS)}  modules={len(counts)}")


if __name__ == "__main__":
    build()
