# -*- coding: utf-8 -*-
"""重建 PGT测试清单模板：业务全量（菜单+自动化），排除若依 system/monitor/tool。不写 bak。"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from hil.report.xlsx_layout import (
    apply_checklist_layout,
    apply_simple_sheet_layout,
    atomic_save,
)

from hil.browser.pages import PAYLOAD_PAGES
from hil.report.catalog import build_catalog

REPO = Path(__file__).resolve().parents[3]
DEFAULT_TEMPLATE = REPO / "doc" / "PGT测试清单模板.xlsx"
SHEET = "功能测试清单"
SUMMARY = "按模块汇总"
DIAG = "映射说明"

HEADERS = [
    "序号", "模块", "页面标题", "测试内容", "测试结果",
    "缺陷编号/备注", "测试人", "测试日期", "自动化ID", "套件",
]

# 若依框架菜单前缀 / 组件：不进业务清单
RUOYI_EXCLUDE_PREFIXES = (
    "/system/",
    "/monitor/",
    "/tool/",
    "/user/profile",
    "/login",
    "/register",
)

# HIL 里仍冒烟但挂在 monitor 下的项：模板不收，回报里问用户
UNCERTAIN_KEEP_OUT = {
    "/monitor/transportCrypto",
}



def _semantic_rows_from_ws(ws) -> list[tuple]:
    """Key columns that define checklist content (ignore result/tester/date)."""
    rows = []
    for r in range(2, (ws.max_row or 1) + 1):
        rows.append(
            (
                r - 1,
                str(ws.cell(r, 2).value or "").strip(),  # module
                str(ws.cell(r, 3).value or "").strip(),  # page
                str(ws.cell(r, 4).value or "").strip(),  # content
                str(ws.cell(r, 9).value or "").strip(),  # auto_id
                str(ws.cell(r, 10).value or "").strip(),  # suite
            )
        )
    return rows


def _template_semantically_equal(path: Path, new_ws) -> bool:
    """True if existing template checklist key cols match newly built sheet."""
    if not path.exists():
        return False
    try:
        from openpyxl import load_workbook
        old = load_workbook(path, data_only=False)
    except Exception:
        return False
    if SHEET not in old.sheetnames:
        return False
    return _semantic_rows_from_ws(old[SHEET]) == _semantic_rows_from_ws(new_ws)

def _is_ruoyi_path(path: str) -> bool:
    p = path if path.startswith("/") else "/" + path
    return any(p.startswith(x) or p == x.rstrip("/") for x in RUOYI_EXCLUDE_PREFIXES)


def _business_pages() -> list[tuple[str, str, str]]:
    """现网业务页 = HIL PAYLOAD_PAGES 去掉若依系统路径。"""
    out = []
    for path, sel, title in PAYLOAD_PAGES:
        if _is_ruoyi_path(path) or path in UNCERTAIN_KEEP_OUT:
            continue
        out.append((path, sel, title))
    return out



def _page_function_rows() -> list[dict]:
    """Deep checklist rows: prefer real auto_id; leftover gaps stay [待深自动化] without id."""
    rows: list[dict] = []
    linked = [
        # Type B passthrough xfer
        ("首页", "首页·dashboard设备服务 /index", "透传开 CAN/串口后调试收发环回（B类）",
         "scenarios/test_hil_passthrough_xfer.py::test_passthrough_can_xfer_loopback", "hil_can"),
        ("调试", "数据收发", "透传串口·调试页收发环回（B类）",
         "scenarios/test_hil_passthrough_xfer.py::test_passthrough_serial_xfer_loopback", "hil_serial"),
        ("调试", "数据收发", "UDP 硬件在环（功能未实现，保留项 skip）",
         "scenarios/test_hil_passthrough_xfer.py::test_udp_hil_feature_not_ready", "hil_udp"),
        # Type A reply placeholders
        ("遥控-BIU", "BIU·遥控回复", "遥控总线回复（未实现/不要求则 skip，保留项）",
         "scenarios/test_hil_passthrough_xfer.py::test_tc_reply_not_implemented[biu]", "hil_can"),
        ("遥控-XL", "XL·遥控回复", "遥控总线回复（未实现/不要求则 skip，保留项）",
         "scenarios/test_hil_passthrough_xfer.py::test_tc_reply_not_implemented[xl]", "hil_can"),
        ("遥控-YJ216", "YJ216·遥控回复", "遥控总线回复（未实现/不要求则 skip，保留项）",
         "scenarios/test_hil_passthrough_xfer.py::test_tc_reply_not_implemented[yj216]", "hil_can"),
        # Covered by existing specialized suites (pointer rows)
        ("调试", "数据模拟", "页面控件+ /payload/telemetry/dev/* 已有 ui_only 专测",
         "scenarios/test_debug_simulate.py::test_simulate_click_start_stop_and_clear", "ui_only"),
        ("调试", "配置文件", "列表/预览/上传/重载已有 ui_only 专测",
         "scenarios/test_debug_config.py::test_config_page_buttons", "ui_only"),
        ("调试", "遥测计算/校验/数转", "golden 对照已有 ui_only 专测",
         "scenarios/test_debug_tmcalc.py::test_tmcalc_field_matches_golden", "ui_only"),
        ("调试", "时间转换", "纪元+周/周内秒/毫秒换算对 Unix ms/ISO UTC，ui_only 专项",
         "scenarios/test_debug_timeconv.py::test_timeconv_protocol_to_unix", "ui_only"),
        ("调试", "时间转换", "日期字符串/Unix 解析对 ISO UTC/Unix ms，ui_only 专项",
         "scenarios/test_debug_timeconv.py::test_timeconv_date_parse", "ui_only"),
        ("遥测", "实时数据", "有实时源时各字段轮询刷新当前值（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_live_fields_poll", "ui_only"),
        ("遥测", "实时曲线", "增加曲线后实时刷新并可导出 CSV（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_live_curve_add_and_export", "ui_only"),
        ("遥测", "历史CAN数据", "选表+时间窗后回放有帧数变化（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_can_history_replay", "ui_only"),
        ("遥测", "历史CAN曲线", "查询历史曲线绘制成功（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_can_archive_curve", "ui_only"),
        ("遥测", "历史文件数据", "上传并选文件后回放可用（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_file_history_replay", "ui_only"),
        ("遥测", "历史文件曲线", "对已解析文件选参画图（接收/显示）",
         "scenarios/test_tm_deep.py::test_tm_file_curve_draw", "ui_only"),
        ("单板", "相机v1.6/v1.7", "图像串口开连接+测试机应答D6出图像素对照（400/其它8倍数）",
         "scenarios/test_hil_camera_image.py::test_camera_image_xfer", "hil_serial"),
        ("单板", "电机控制板", "开 UDP 后工程遥测字段刷新（接收/显示，业务未落地则 skip）",
         "scenarios/test_tm_deep.py::test_board_dj_udp_tm_pending", "ui_only"),
    ]
    for mod, page, content, aid, suite in linked:
        rows.append({
            "module": mod,
            "page": page,
            "content": content,
            "auto_id": aid,
            "suite": suite,
        })
    # Remaining deep gaps (no auto yet)
    pending = [
        # formerly blank rows now linked above (camera image / tm deep / dj skip)
    ]
    for mod, page, content in pending:
        rows.append({
            "module": mod,
            "page": page,
            "content": f"[待深自动化] {content}",
            "auto_id": "",
            "suite": "manual_page",
        })
    return rows


def _filter_catalog(catalog: list[dict]) -> list[dict]:
    """自动化目录：丢掉明显若依系统页的 UI 冒烟；HIL/unit/db 全留。"""
    out = []
    for it in catalog:
        nid = it.get("nodeid") or it.get("auto_id") or ""
        page = it.get("page") or ""
        # ui page_opens / key_element for monitor/system/tool
        if "test_ui_pages.py" in nid:
            # extract path from nodeid [...] 
            if "/monitor/" in nid or "/system/" in nid or "/tool/" in nid:
                continue
            if "/monitor/" in page or "系统监控" in (it.get("module") or ""):
                # transportCrypto 等
                continue
        out.append(it)
    return out


def sync_template(template: Path | None = None, keep_manual: bool = True) -> dict:
    path = Path(template) if template else DEFAULT_TEMPLATE
    path.parent.mkdir(parents=True, exist_ok=True)

    catalog = _filter_catalog(build_catalog(include_db_placeholders=True))
    page_rows = _page_function_rows() if keep_manual else []

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for col, name in enumerate(HEADERS, 1):
        ws.cell(1, col, name).font = Font(bold=True)

    r = 2
    for it in catalog:
        ws.cell(r, 1, r - 1)
        ws.cell(r, 2, it["module"])
        ws.cell(r, 3, it["page"])
        ws.cell(r, 4, it["content"])
        ws.cell(r, 9, it["auto_id"])
        ws.cell(r, 10, it["suite"])
        for c in range(1, 11):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    auto_count = len(catalog)

    for it in page_rows:
        ws.cell(r, 1, r - 1)
        ws.cell(r, 2, it["module"])
        ws.cell(r, 3, it["page"])
        ws.cell(r, 4, it["content"])
        ws.cell(r, 9, it.get("auto_id") or None)
        ws.cell(r, 10, it.get("suite") or "manual_page")
        for c in range(1, 11):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    mods = Counter()
    for rr in range(2, r):
        mods[str(ws.cell(rr, 2).value or "未分类")] += 1

    ws2 = wb.create_sheet(SUMMARY)
    for c, h in enumerate(["模块", "用例数", "建议测试人", "完成数", "通过数", "备注"], 1):
        ws2.cell(1, c, h).font = Font(bold=True)
    rr = 2
    for mod, n in sorted(mods.items(), key=lambda x: (-x[1], x[0])):
        ws2.cell(rr, 1, mod)
        ws2.cell(rr, 2, n)
        rr += 1
    ws2.cell(rr, 1, "合计")
    ws2.cell(rr, 2, sum(mods.values()))

    ws3 = wb.create_sheet(DIAG)
    ws3.cell(1, 1, "说明").font = Font(bold=True)
    ws3.cell(
        2,
        1,
        "由 python -m hil.report sync 重建：以 pytest 实采自动化 + 现网业务页（排除若依 system/monitor/tool）为准；"
        "旧清单不作权威。自动化ID=pytest nodeid；空=页面功能待补自动化。填报：python -m hil.report fill <run_dir>。"
        "禁止生成 .bak。",
    )

    apply_checklist_layout(ws)
    apply_simple_sheet_layout(ws2)
    apply_simple_sheet_layout(ws3, {1: 100})

    skipped_write = False
    if _template_semantically_equal(path, ws):
        skipped_write = True
    else:
        atomic_save(wb, path)

    present = {
        str(ws.cell(rr, 9).value).strip()
        for rr in range(2, ws.max_row + 1)
        if ws.cell(rr, 9).value
    }
    missing = [it["auto_id"] for it in catalog if it["auto_id"] not in present]

    return {
        "template": str(path),
        "backup": None,
        "rows_total": ws.max_row - 1,
        "auto_rows": auto_count,
        "page_rows": len(page_rows),
        "catalog": len(catalog),
        "missing_auto_ids": missing,
        "modules": dict(mods),
        "skipped_write": skipped_write,
        "written": not skipped_write,
        "excluded_ruoyi_note": "system/monitor/tool；transportCrypto 按结论保持不测",
    }


if __name__ == "__main__":
    import json

    print(json.dumps(sync_template(), ensure_ascii=False, indent=2))
