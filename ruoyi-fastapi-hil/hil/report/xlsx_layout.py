# -*- coding: utf-8 -*-
"""Excel 清单排版：列宽 / 行高 / 表头，供 sync 与 fill 共用。"""
from __future__ import annotations

from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 功能测试清单 10 列：序号..套件
CHECKLIST_WIDTHS = {
    1: 6,
    2: 14,
    3: 42,
    4: 56,
    5: 10,
    6: 28,
    7: 12,
    8: 12,
    9: 52,
    10: 14,
}

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True, name="微软雅黑", size=11)
CELL_FONT = Font(name="微软雅黑", size=10)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")

MAX_ROW_HEIGHT = 48.0
DEFAULT_ROW_HEIGHT = 18.0
HEADER_ROW_HEIGHT = 22.0


def _estimate_lines(text: str, col_width: float) -> int:
    if not text:
        return 1
    s = str(text)
    width = max(int(col_width), 4)
    lines = 0
    for part in s.splitlines() or [""]:
        # 中文按全角近似：openpyxl 列宽约等于拉丁字符数
        n = 0
        for ch in part:
            n += 2 if ord(ch) > 0x7F else 1
        lines += max(1, (n + width - 1) // width)
    return max(1, lines)


def apply_checklist_layout(ws) -> None:
    """主清单表：列宽、表头样式、适中行高、冻结首行。"""
    for col, w in CHECKLIST_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    max_col = max(CHECKLIST_WIDTHS)
    for c in range(1, max_col + 1):
        cell = ws.cell(1, c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN

    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max(ws.max_row, 1)}"

    for r in range(2, ws.max_row + 1):
        max_lines = 1
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
            w = CHECKLIST_WIDTHS.get(c, 12)
            max_lines = max(max_lines, _estimate_lines(cell.value, w))
        h = min(MAX_ROW_HEIGHT, max(DEFAULT_ROW_HEIGHT, 14.0 * max_lines))
        ws.row_dimensions[r].height = h


def apply_simple_sheet_layout(ws, widths: dict[int, float] | None = None) -> None:
    """汇总/说明等附属表的轻量排版。"""
    widths = widths or {1: 24, 2: 14, 3: 14, 4: 14, 5: 14, 6: 40}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    if ws.max_row >= 1:
        for c in range(1, max(widths) + 1):
            cell = ws.cell(1, c)
            if cell.value is not None:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = HEADER_ALIGN
        ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = DEFAULT_ROW_HEIGHT
        for c in range(1, max(widths) + 1):
            cell = ws.cell(r, c)
            if cell.value is not None:
                cell.font = CELL_FONT
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def atomic_save(wb, path) -> None:
    """先写 .xlsx.tmp 再 replace，避免中断留下坏 zip。"""
    path = Path(path)
    tmp = path.with_suffix(path.suffix + ".tmp")
    if tmp.exists():
        tmp.unlink()
    wb.save(tmp)
    tmp.replace(path)
